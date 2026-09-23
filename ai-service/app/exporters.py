from __future__ import annotations

from io import BytesIO
from pathlib import Path

from app.errors import ServiceError
from app.schemas import ProcessResponse


def export_docx(protocol: ProcessResponse) -> bytes:
    try:
        from docx import Document
        from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Mm, Pt, RGBColor
    except ImportError as exc:
        raise ServiceError(
            "EXPORT_DEPENDENCY_MISSING", "python-docx is not installed", status_code=503
        ) from exc

    document = Document()
    section = document.sections[0]
    section.top_margin = Mm(15)
    section.bottom_margin = Mm(15)
    section.left_margin = Mm(15)
    section.right_margin = Mm(15)
    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)
    styles["Normal"].paragraph_format.space_after = Pt(4)
    for style_name, size in (("Title", 20), ("Heading 1", 14)):
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.underline = False
        borders = style.element.get_or_add_pPr().find(qn("w:pBdr"))
        if borders is not None:
            style.element.get_or_add_pPr().remove(borders)
    document.add_heading(f"Протокол совещания {protocol.meeting_id}", level=0)
    document.add_heading("Краткое содержание", level=1)
    document.add_paragraph(protocol.summary or "—")
    document.add_heading("Поручения", level=1)
    if protocol.action_items:
        table = document.add_table(rows=1, cols=5)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        column_widths = [Mm(10), Mm(73), Mm(37), Mm(30), Mm(30)]
        for column, width in zip(table.columns, column_widths, strict=True):
            column.width = width
        headers = ["№", "Поручение", "Ответственный", "Срок", "Основание"]
        for cell, header, width in zip(
            table.rows[0].cells, headers, column_widths, strict=True
        ):
            cell.text = header
            cell.width = width
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for run in cell.paragraphs[0].runs:
                run.bold = True
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "E8EEF7")
            cell._tc.get_or_add_tcPr().append(shading)
        for item in protocol.action_items:
            cells = table.add_row().cells
            values = [
                str(item.id + 1),
                item.text,
                item.assignee_participant_id or item.assignee_speaker_label or "Не указан",
                item.due_date.isoformat() if item.due_date else (item.due_text_raw or "Не указан"),
                ", ".join(map(str, item.evidence_segment_ids)) or "—",
            ]
            for cell, value, width in zip(cells, values, column_widths, strict=True):
                cell.text = value
                cell.width = width
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    else:
        document.add_paragraph("Поручения не зафиксированы.")
    document.add_heading("Транскрипт", level=1)
    for segment in protocol.segments:
        stamp = _timestamp(segment.start_ms)
        paragraph = document.add_paragraph()
        paragraph.add_run(f"[{stamp}] {segment.speaker}: ").bold = True
        paragraph.add_run(segment.text)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def export_pdf(protocol: ProcessResponse) -> bytes:
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_CENTER
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    except ImportError as exc:
        raise ServiceError(
            "EXPORT_DEPENDENCY_MISSING", "reportlab is not installed", status_code=503
        ) from exc

    font = _register_unicode_font(pdfmetrics, TTFont)
    output = BytesIO()
    doc = SimpleDocTemplate(output, pagesize=A4, rightMargin=15 * mm, leftMargin=15 * mm)
    styles = getSampleStyleSheet()
    body = ParagraphStyle(
        "BodyUnicode", parent=styles["BodyText"], fontName=font, fontSize=9, leading=12
    )
    heading = ParagraphStyle(
        "HeadingUnicode", parent=styles["Heading2"], fontName=font, spaceBefore=8
    )
    title = ParagraphStyle(
        "TitleUnicode", parent=styles["Title"], fontName=font, alignment=TA_CENTER
    )
    story = [Paragraph(_xml(f"Протокол совещания {protocol.meeting_id}"), title), Spacer(1, 6 * mm)]
    story.extend(
        [Paragraph("Краткое содержание", heading), Paragraph(_xml(protocol.summary or "—"), body)]
    )
    story.append(Paragraph("Поручения", heading))
    rows = [[Paragraph(_xml(value), body) for value in ["№", "Поручение", "Ответственный", "Срок"]]]
    for item in protocol.action_items:
        rows.append(
            [
                Paragraph(str(item.id + 1), body),
                Paragraph(_xml(item.text), body),
                Paragraph(
                    _xml(
                        item.assignee_participant_id or item.assignee_speaker_label or "Не указан"
                    ),
                    body,
                ),
                Paragraph(
                    _xml(
                        item.due_date.isoformat()
                        if item.due_date
                        else (item.due_text_raw or "Не указан")
                    ),
                    body,
                ),
            ]
        )
    if len(rows) == 1:
        rows.append(["—", Paragraph("Поручения не зафиксированы", body), "—", "—"])
    table = Table(rows, colWidths=[10 * mm, 88 * mm, 42 * mm, 30 * mm], repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, -1), font),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EEF7")),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    story.extend([table, Paragraph("Транскрипт", heading)])
    for segment in protocol.segments:
        story.append(
            Paragraph(
                _xml(f"[{_timestamp(segment.start_ms)}] {segment.speaker}: {segment.text}"), body
            )
        )
    doc.build(story)
    return output.getvalue()


def _register_unicode_font(pdfmetrics, ttfont) -> str:
    candidates = [
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        Path("C:/Windows/Fonts/arial.ttf"),
        Path("C:/Windows/Fonts/calibri.ttf"),
    ]
    for path in candidates:
        if path.exists():
            name = "ProtocolUnicode"
            if name not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(ttfont(name, str(path)))
            return name
    raise ServiceError(
        "EXPORT_FONT_MISSING", "a Unicode TTF font is required for PDF export", status_code=503
    )


def _timestamp(milliseconds: int) -> str:
    seconds = milliseconds // 1000
    return f"{seconds // 60:02d}:{seconds % 60:02d}"


def _xml(value: str) -> str:
    from xml.sax.saxutils import escape

    return escape(value).replace("\n", "<br/>")
